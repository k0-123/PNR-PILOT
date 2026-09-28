"""Step B: captured website pages -> result columns -> final Excel (worker + page_results)."""
import re
from types import SimpleNamespace

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import lookups
from app.api import create_app
from app.core.models import JobStatus, LookupStatus
from app.extraction.gemini_client import GeminiClient
from app.jobs import reparse_all, retry_failed, start_lookups
from app.lookup import page_results
from app.worker import Worker
from tests.conftest import FakeModels, make_job

LS = LookupStatus
PAGE = "Your booking reference: {pnr}\nLast name: {surname}\n" + "Flight MH 144 AKL-KUL 27 SEP 2026 CONFIRMED\n" * 3


def pax(surname, first, **values):
    base = {"surname": surname, "first_name": first, "flight_numbers": "MH 144", "route": "AKL-KUL",
            "travel_dates": "27SEP", "booking_status": "CONFIRMED", "ticket_numbers": None,
            "passenger_names": f"{first} {surname}", "confidence": 0.95}
    return {**base, **values}


def booking(pnr, *passengers):
    return {"bookings": [{"pnr": pnr, "pnr_confidence": 0.99, "passengers": list(passengers)}]}


class Pages:
    """Fake Gemini for page reading: answers by the PNR named in the prompt."""

    def __init__(self, answers: dict):
        self.answers = answers
        self.seen: list[str] = []

    def __call__(self, contents):
        pnr = re.search(r'searching booking reference "([A-Z0-9]+)"', contents[0]).group(1)
        self.seen.append(pnr)
        answer = self.answers[pnr]
        if isinstance(answer, list):  # a sequence of answers for repeated reads
            return answer.pop(0)
        return answer


def worker_for(settings, answers):
    pages = Pages(answers)
    fake = FakeModels(pages)
    factory = lambda: GeminiClient(settings, client=SimpleNamespace(models=fake), sleep=lambda s: None)  # noqa: E731
    return Worker(settings, gemini_factory=factory), pages


def captured_job(db, storage, settings, rows, family=False):
    """A job opened for lookups whose PNRs were all captured by the extension."""
    job = make_job(db, rows)
    if family:  # row 2 = second passenger (other first name) on row 1's booking
        db.conn.execute("UPDATE rows SET first_name='RAJ', extraction_status='READY', duplicate_of=NULL "
                        "WHERE job_id=? AND seq=1", (job,))
        db.conn.commit()
    start_lookups(db, job)
    token = lookups.authenticate(db, lookups.create_token(db, "Asha"))
    lookups.claim(db, job, token, 50, settings)
    for r in {r["effective_pnr"]: r for r in db.get_rows(job)}.values():
        lookups.save_capture(db, storage, settings, job, r["effective_pnr"], token,
                             text=PAGE.format(pnr=r["effective_pnr"], surname=r["effective_surname"]),
                             screenshot_b64=None, url="https://online.example/booking/manage-booking/confirmation")
    return job, token


def results(db, job):
    import json
    return [(r["effective_surname"], r["effective_first_name"], r["lookup_status"],
             json.loads(r["result_json"]) if r["result_json"] else None, r["error_message"])
            for r in db.get_rows(job)]


def lookup(db, job, pnr):
    return db.conn.execute("SELECT * FROM lookups WHERE job_id=? AND pnr=?", (job, pnr)).fetchone()


# ------------------------------------------------------------------------------------------

def test_family_booking_fills_each_passengers_row_and_job_completes(settings, db, storage):
    job, _ = captured_job(db, storage, settings, [("SHUKLA", "FAM001"), ("SHUKLA", "FAM001"), ("PATEL", "OKB002")],
                          family=True)
    worker, pages = worker_for(settings, {
        "FAM001": booking("FAM001", pax("SHUKLA", "SIDDHESHVA", ticket_numbers="232-111"),
                          pax("SHUKLA", "RAJ", ticket_numbers="232-222")),
        "OKB002": booking("OKB002", pax("PATEL", "SIDDHESHVA", ticket_numbers="232-333")),
    })
    assert worker.run_once()
    assert sorted(pages.seen) == ["FAM001", "OKB002"]  # one read per PNR, not per row
    got = results(db, job)
    assert [(s, f, st) for s, f, st, _, _ in got] == [("SHUKLA", "SIDDHESHVA", LS.PARSED),
                                                     ("SHUKLA", "RAJ", LS.PARSED), ("PATEL", "SIDDHESHVA", LS.PARSED)]
    assert [v["ticket_numbers"] for _, _, _, v, _ in got] == ["232-111", "232-222", "232-333"]  # per passenger
    assert all(v["flight_numbers"] == "MH 144" for _, _, _, v, _ in got)  # booking-wide
    assert lookup(db, job, "FAM001")["status"] == LS.PARSED

    worker.run_once()  # housekeeping: finish the job + final Excel
    assert db.get_job(job)["status"] == JobStatus.COMPLETED
    ws = load_workbook(storage.path(f"outputs/job_{job}/job_{job}_final.xlsx"))["Results"]
    headers = [c.value for c in ws[1]]
    vals = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    col = {h: i for i, h in enumerate(headers)}
    assert [v[col["Ticket Number(s)"]] for v in vals] == ["232-111", "232-222", "232-333"]
    assert {v[col["Looked Up By"]] for v in vals} == {"Asha"}
    assert all("manage-booking/confirmation" in v[col["Source"]] for v in vals)


def test_passenger_missing_from_page_gets_booking_details_only(settings, db, storage):
    job, _ = captured_job(db, storage, settings, [("SHUKLA", "FAM001"), ("SHUKLA", "FAM001")], family=True)
    worker, _ = worker_for(settings, {
        "FAM001": booking("FAM001", pax("SHUKLA", "SIDDHESHVA", ticket_numbers="232-111"))})
    worker.run_once()
    (_, _, st1, v1, _), (_, _, st2, v2, note2) = results(db, job)
    assert (st1, v1["ticket_numbers"]) == (LS.PARSED, "232-111")
    assert st2 == LS.PARSED and "ticket_numbers" not in v2 and v2["flight_numbers"] == "MH 144"
    assert "booking details only" in note2


def test_page_for_another_booking_is_a_mismatch(settings, db, storage):
    job, _ = captured_job(db, storage, settings, [("SHUKLA", "OKC001")])
    worker, _ = worker_for(settings, {"OKC001": booking("ZZZ999", pax("OTHER", "X"))})
    worker.run_once()
    ((_, _, st, values, note),) = results(db, job)
    assert (st, values) == (LS.MISMATCH, None) and "ZZZ999" in note


def test_unreadable_page_is_read_again_on_retry_without_a_new_search(settings, db, storage):
    job, _ = captured_job(db, storage, settings, [("SHUKLA", "OKD001"), ("PATEL", "OKD002")])
    worker, pages = worker_for(settings, {
        "OKD001": [{"bookings": []}, booking("OKD001", pax("SHUKLA", "SIDDHESHVA"))],
        "OKD002": booking("OKD002", pax("PATEL", "SIDDHESHVA")),
    })
    worker.run_once()
    assert [st for _, _, st, _, _ in results(db, job)] == [LS.PARSE_ERROR, LS.PARSED]
    worker.run_once()
    assert db.get_job(job)["status"] == JobStatus.COMPLETED_WITH_ERRORS
    captures_before = lookup(db, job, "OKD001")["captures"]

    assert retry_failed(db, job) == 1
    assert lookup(db, job, "OKD001")["status"] == LS.CAPTURED  # re-read, not back to the website queue
    worker.run_once()
    worker.run_once()
    assert [st for _, _, st, _, _ in results(db, job)] == [LS.PARSED, LS.PARSED]
    assert lookup(db, job, "OKD001")["captures"] == captures_before  # no new website search
    assert db.get_job(job)["status"] == JobStatus.COMPLETED
    assert pages.seen.count("OKD001") == 2


def test_gemini_failure_is_a_parse_error(settings, db, storage):
    job, _ = captured_job(db, storage, settings, [("SHUKLA", "OKE001")])
    worker, _ = worker_for(settings, {"OKE001": "not json"})
    worker.run_once()
    ((_, _, st, _, note),) = results(db, job)
    assert st == LS.PARSE_ERROR and "could not read the page" in note


def test_recapture_while_reading_discards_the_old_result(settings, db, storage):
    job, token = captured_job(db, storage, settings, [("SHUKLA", "OKF001")])
    (cap,) = page_results.claim(db, "w1", 5)
    cap = page_results.prepare(db, storage, cap)
    lookups.save_capture(db, storage, settings, job, "OKF001", token, text=PAGE.format(pnr="OKF001", surname="X"),
                         screenshot_b64=None, url=None, recapture=True)
    fake = SimpleNamespace(data=SimpleNamespace(bookings=[]), cost_usd=0.0)
    assert page_results.apply(db, settings, None, cap, fake, []) == LS.CAPTURED  # stale: ignored
    assert lookup(db, job, "OKF001")["status"] == LS.CAPTURED


def test_two_workers_never_read_the_same_page(settings, db, storage):
    captured_job(db, storage, settings, [("A", "OKG001"), ("B", "OKG002"), ("C", "OKG003")])
    first = page_results.claim(db, "w1", 2)
    second = page_results.claim(db, "w2", 5)
    assert len(first) == 2 and len(second) == 1
    assert not {c.pnr for c in first} & {c.pnr for c in second}


def test_reparse_all_reads_saved_pages_again(settings, db, storage):
    job, _ = captured_job(db, storage, settings, [("SHUKLA", "OKH001")])
    worker, pages = worker_for(settings, {"OKH001": [booking("OKH001", pax("SHUKLA", "SIDDHESHVA", route="A-B")),
                                                     booking("OKH001", pax("SHUKLA", "SIDDHESHVA", route="C-D"))]})
    worker.run_once()
    worker.run_once()
    assert db.get_job(job)["status"] == JobStatus.COMPLETED
    assert reparse_all(db, job) == 1
    assert db.get_job(job)["status"] == JobStatus.LOOKING_UP
    worker.run_once()
    worker.run_once()
    ((_, _, st, values, _),) = results(db, job)
    assert (st, values["route"], db.get_job(job)["status"]) == (LS.PARSED, "C-D", JobStatus.COMPLETED)


def test_job_waits_while_bookings_are_still_open(settings, db, storage):
    job = make_job(db, [("A", "OKI001"), ("B", "OKI002")])
    start_lookups(db, job)
    token = lookups.authenticate(db, lookups.create_token(db, "Asha"))
    lookups.claim(db, job, token, 1, settings)
    lookups.save_capture(db, storage, settings, job, "OKI001", token, text=PAGE.format(pnr="OKI001", surname="A"),
                         screenshot_b64=None, url=None)
    worker, _ = worker_for(settings, {"OKI001": booking("OKI001", pax("A", "SIDDHESHVA"))})
    worker.run_once()
    worker.run_once()
    assert db.get_job(job)["status"] == JobStatus.LOOKING_UP  # OKI002 not searched yet
    assert storage.exists(f"outputs/job_{job}/job_{job}_final.xlsx")  # live Excel already there


def test_final_excel_download_endpoint(settings, db, storage):
    job, _ = captured_job(db, storage, settings, [("SHUKLA", "OKJ001")])
    token = lookups.create_token(db, "Ben")
    client = TestClient(create_app(settings))
    r = client.get(f"/api/jobs/{job}/final.xlsx", headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 200 and r.content[:2] == b"PK"  # an xlsx (zip) file
    assert client.get(f"/api/jobs/{job}/final.xlsx").status_code == 401


def test_split_glued_title():
    from app.lookup.page_results import split_glued_title
    assert split_glued_title("MsLijing Chen") == "Ms Lijing Chen"
    assert split_glued_title("MrsAnna Rao") == "Mrs Anna Rao"
    assert split_glued_title("MS LIJING CHEN") == "MS LIJING CHEN"   # already fine
    assert split_glued_title("Msumba Joseph") == "Msumba Joseph"     # a name, not a title
    assert split_glued_title("Drake Bell") == "Drake Bell"
    assert split_glued_title(None) is None


def test_redact_personal_removes_passport_and_birth_data():
    from app.lookup.redact import redact_personal
    page = "\n".join([
        "MsLijing Chen", "E-Ticket numbers: 232-2485324958", "Gender: Female",
        "Date of birth: 13/08/1993", "Nationality: Chinese", "Passport number: EG1168320",
        "Passport expiry: 21/04/2029", "Issuing country: China",
        "Date of birth", "01/02/1990",  # label and value on separate lines
        "Passport No.X1234567",         # value glued to the label
        "Primary contact details", "+60 128925888",
    ])
    out = redact_personal(page)
    for secret in ("13/08/1993", "Chinese", "EG1168320", "21/04/2029", "China", "01/02/1990", "X1234567"):
        assert secret not in out, secret
    for kept in ("232-2485324958", "Gender: Female", "+60 128925888", "Primary contact details"):
        assert kept in out, kept
    assert "Passport number: [removed]" in out
