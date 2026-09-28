"""Booking details from GDS PNR display screenshots (fake Gemini)."""
import json
from types import SimpleNamespace

from google import genai
from google.genai import _transformers
from openpyxl import load_workbook

from app.core.config import load_result_fields
from app.core.models import ExtractionStatus, JobStatus, LookupStatus
from app.extraction.gemini_client import GeminiClient
from app.jobs import add_pnr_screens
from app.lookup.gds_screens import build_screen_model, match_passenger
from app.worker import Worker
from tests.conftest import FakeModels, make_job, png


def pax(surname, first="X", **fields):
    base = {"surname": surname, "first_name": first, "flight_numbers": "MH 146", "route": "AKL-KUL",
            "travel_dates": "04SEP", "booking_status": "HK", "ticket_numbers": None,
            "passenger_names": None, "confidence": 0.95}
    return {**base, **fields}


def worker_with(settings, bookings):
    fake = FakeModels(lambda contents: {"bookings": bookings})
    factory = lambda: GeminiClient(settings, client=SimpleNamespace(models=fake), sleep=lambda s: None)  # noqa: E731
    return Worker(settings, gemini_factory=factory), fake


def test_screen_schema_is_generated_from_fields_and_converts():
    model = build_screen_model(load_result_fields())
    schema = _transformers.t_schema(genai.Client(api_key="x")._api_client, model)
    pax_props = schema.properties["bookings"].items.properties["passengers"].items.properties
    assert {"surname", "first_name", "flight_numbers", "ticket_numbers", "confidence"} <= set(pax_props)


def test_match_passenger_rules():
    rows = [{"effective_surname": "KISHTAPPA NAIDU"}, {"effective_surname": "HARIDOSS"}]
    assert match_passenger(rows, "KISHTAPPANAIDU") == [rows[0]]      # spaces ignored
    assert match_passenger(rows, "HARIDOSSKUMAR") == [rows[1]]       # list screen truncated the name
    assert match_passenger(rows, "SMITH") == []
    assert match_passenger(rows, None) == []


def test_worker_fills_rows_from_screen_and_completes_job(settings, db, storage):
    job = make_job(db, [("CHEN", "ER7P5B"), ("HARIDOSS", "9P4BIE"), ("KISHTAPPA NAIDU", "9P4BIE")],
                   status=JobStatus.READY_FOR_LOOKUP)
    worker, fake = worker_with(settings, [
        {"pnr": "ER7P5B", "pnr_confidence": 0.99, "passengers": [pax("CHEN", "LIJING", ticket_numbers="232-1")]},
        {"pnr": "9P4BIE", "pnr_confidence": 0.99, "passengers": [
            pax("HARIDOSS", "MANJULA", ticket_numbers="232-2"),
            pax("KISHTAPPANAIDU", "H", ticket_numbers="232-3", confidence=0.5)]},
    ])
    add_pnr_screens(db, storage, settings, job, [("rt.png", png(11))])
    assert worker.run_once()

    rows = db.get_rows(job)
    assert [r["lookup_status"] for r in rows] == [LookupStatus.PARSED] * 3
    assert [json.loads(r["result_json"])["ticket_numbers"] for r in rows] == ["232-1", "232-2", "232-3"]
    assert "low confidence" in rows[2]["error_message"]            # flagged for a human check
    assert rows[0]["screenshot_path"].startswith(f"screens/job_{job}/")
    screen = db.get_pnr_screens(job)[0]
    assert (screen["status"], screen["pnrs"], screen["matched_rows"]) == ("DONE", "9P4BIE, ER7P5B", 3)
    assert db.get_job(job)["status"] == JobStatus.COMPLETED
    ws = load_workbook(storage.path(f"outputs/job_{job}/job_{job}_final.xlsx"))["Results"]
    headers = [c.value for c in ws[1]]
    assert ws.cell(2, headers.index("Ticket Number(s)") + 1).value == "232-1"
    # the extraction model (Pro) reads screens: exact characters matter
    assert fake.calls[0]["model"] == settings.gemini_model_extract


def test_unmatched_passengers_and_partial_job(settings, db, storage):
    job = make_job(db, [("CHEN", "ER7P5B"), ("PATEL", "OKX777")], status=JobStatus.READY_FOR_LOOKUP)
    worker, _ = worker_with(settings, [
        {"pnr": "ER7P5B", "pnr_confidence": 0.99, "passengers": [pax("CHEN"), pax("WONG")]},
    ])
    add_pnr_screens(db, storage, settings, job, [("rt.png", png(12))])
    worker.run_once()
    screen = db.get_pnr_screens(job)[0]
    assert screen["matched_rows"] == 1 and json.loads(screen["unmatched"]) == ["ER7P5B WONG/X"]
    assert [r["lookup_status"] for r in db.get_rows(job)] == [LookupStatus.PARSED, LookupStatus.NOT_STARTED]
    assert db.get_job(job)["status"] == JobStatus.READY_FOR_LOOKUP  # PATEL still has no result


def test_single_passenger_matches_on_pnr_even_if_surname_misread(settings, db, storage):
    job = make_job(db, [("CHEN", "ER7P5B")], status=JobStatus.READY_FOR_LOOKUP)
    worker, _ = worker_with(settings, [{"pnr": "ER7P5B", "passengers": [pax("CHFN")]}])
    add_pnr_screens(db, storage, settings, job, [("rt.png", png(13))])
    worker.run_once()
    assert db.get_rows(job)[0]["lookup_status"] == LookupStatus.PARSED


def test_rows_needing_review_are_not_filled(settings, db, storage):
    job = make_job(db, [("CHEN", "ER7P5B")], status=JobStatus.AWAITING_REVIEW)
    db.conn.execute("UPDATE rows SET extraction_status='NEEDS_REVIEW' WHERE job_id=?", (job,))
    db.conn.commit()
    worker, _ = worker_with(settings, [{"pnr": "ER7P5B", "passengers": [pax("CHEN")]}])
    add_pnr_screens(db, storage, settings, job, [("rt.png", png(14))])
    worker.run_once()
    screen = db.get_pnr_screens(job)[0]
    assert screen["status"] == "FAILED" and "no passenger matched" in screen["error"]
    assert db.get_rows(job)[0]["extraction_status"] == ExtractionStatus.NEEDS_REVIEW


def test_screen_without_pnr_display_fails_clearly(settings, db, storage):
    job = make_job(db, [("CHEN", "ER7P5B")], status=JobStatus.READY_FOR_LOOKUP)
    worker, _ = worker_with(settings, [])
    add_pnr_screens(db, storage, settings, job, [("cat.png", png(15))])
    worker.run_once()
    screen = db.get_pnr_screens(job)[0]
    assert screen["status"] == "FAILED" and "no PNR display" in screen["error"]


def test_same_screen_uploaded_twice_is_not_duplicated(settings, db, storage):
    job = make_job(db, [("CHEN", "ER7P5B")], status=JobStatus.READY_FOR_LOOKUP)
    add_pnr_screens(db, storage, settings, job, [("a.png", png(16)), ("b.png", png(16))])
    assert len(db.get_pnr_screens(job)) == 1


def test_identical_screen_is_served_from_cache_without_a_second_gemini_call(settings, db, storage):
    # Job A reads the screen for real.
    job_a = make_job(db, [("CHEN", "ER7P5B")], status=JobStatus.READY_FOR_LOOKUP)
    worker_a, fake_a = worker_with(settings, [
        {"pnr": "ER7P5B", "pnr_confidence": 0.99, "passengers": [pax("CHEN", "LIJING", ticket_numbers="232-1")]},
    ])
    add_pnr_screens(db, storage, settings, job_a, [("rt.png", png(21))])
    worker_a.run_once()
    assert len(fake_a.calls) == 1
    assert db.get_pnr_screens(job_a)[0]["status"] == "DONE"

    # Job B uploads the SAME screenshot bytes (same SHA-256). The read is served from job A's cached
    # response, so Gemini is not called and the CACHED values (not this worker's answer) fill the row.
    job_b = make_job(db, [("CHEN", "ER7P5B")], status=JobStatus.READY_FOR_LOOKUP)
    worker_b, fake_b = worker_with(settings, [
        {"pnr": "ER7P5B", "passengers": [pax("CHEN", "LIJING", ticket_numbers="999-SHOULD-NOT-BE-USED")]},
    ])
    add_pnr_screens(db, storage, settings, job_b, [("rt.png", png(21))])
    worker_b.run_once()
    assert fake_b.calls == []  # nothing sent to Gemini: reused the cache
    row = db.get_rows(job_b)[0]
    assert row["lookup_status"] == LookupStatus.PARSED
    assert json.loads(row["result_json"])["ticket_numbers"] == "232-1"
    screen_b = db.get_pnr_screens(job_b)[0]
    assert (screen_b["status"], screen_b["cost_usd"]) == ("DONE", 0.0)


def test_family_with_same_surname_matched_one_to_one(settings, db, storage):
    job = make_job(db, [("TESTPAX", "TQ4X8Z"), ("TESTPAX", "TQ4X8Z")], status=JobStatus.READY_FOR_LOOKUP)
    ids = [r["id"] for r in db.get_rows(job)]
    for rid, first in zip(ids, ["ALICE", "BOB"]):  # distinct passengers, same surname + PNR
        db.conn.execute("UPDATE rows SET first_name=?, extraction_status='READY', duplicate_of=NULL WHERE id=?",
                        (first, rid))
    db.conn.commit()
    # screen lists Bob first: matching must follow first names, not order
    worker, _ = worker_with(settings, [{"pnr": "TQ4X8Z", "passengers": [
        pax("TESTPAX", "BOB", ticket_numbers="232-B"), pax("TESTPAX", "ALICE", ticket_numbers="232-A")]}])
    add_pnr_screens(db, storage, settings, job, [("rt.png", png(17))])
    worker.run_once()
    tickets = [json.loads(r["result_json"])["ticket_numbers"] for r in db.get_rows(job)]
    assert tickets == ["232-A", "232-B"]
    assert db.get_pnr_screens(job)[0]["matched_rows"] == 2
