"""The final Excel/CSV."""
from openpyxl import load_workbook

from app.core.config import load_result_fields
from app.core.models import LookupStatus
from app.jobs import export_final
from tests.conftest import make_job

S = LookupStatus


# ---- final Excel ----

def test_final_excel_and_csv(db, storage, settings):
    job = make_job(db, [("SHUKLA", "OKP001"), ("B", "NFP002"), ("C", "ERR5P3"), ("=EVIL", "OKP004")])
    rows = db.get_rows(job)
    db.save_lookup_result(rows[0]["id"], result={"flight_numbers": "AI 001"}, page_text_path=None,
                          screenshot_path=None)
    for row, status in zip(rows, [S.PARSED, S.NOT_FOUND, S.MISMATCH]):
        db.finish_row(row["id"], status, attempts=1)

    xlsx_key, csv_key = export_final(db, storage, job)
    assert xlsx_key.endswith(f"job_{job}_final.xlsx")

    wb = load_workbook(storage.path(xlsx_key))
    ws = wb["Results"]
    headers = [c.value for c in ws[1]]
    labels = [f.label for f in load_result_fields().fields]
    # who the row is, then what the lookup found (on screen when the file opens), then the rest
    assert headers[:7] == ["Image name", "Line No", "Pax", "Surname", "First Name", "Title", "PNR"]
    assert headers[7:7 + len(labels)] == labels
    assert headers[7 + len(labels):9 + len(labels)] == ["Lookup Status", "Error"]
    assert headers[9 + len(labels)] == "Class"
    assert headers[-6:] == ["Edited Surname", "Edited PNR", "Attempts", "Looked Up At", "Looked Up By", "Source"]
    vals = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    col = {h: i for i, h in enumerate(headers)}
    # the invalid "=EVIL" row needs review: never looked up, so no lookup status
    assert [v[col["Lookup Status"]] for v in vals] == ["PARSED", "NOT_FOUND", "MISMATCH", None]
    evil = ws.cell(row=5, column=col["Surname"] + 1)
    assert (evil.value, evil.data_type, evil.quotePrefix) == ("=EVIL", "s", True)  # text, never a formula
    assert vals[0][col["Flight Number(s)"]] == "AI 001"

    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=2)}
    assert (summary["Total rows"], summary["Success"], summary["Not found"], summary["Failed"],
            summary["Needs review"]) == (4, 1, 1, 1, 1)
    assert "Total Gemini cost (USD)" in summary and "Job duration" in summary

    csv_text = storage.read_bytes(csv_key).decode("utf-8-sig")
    assert csv_text.splitlines()[0].startswith("Image name,Line No")
    assert "'=EVIL" in csv_text
