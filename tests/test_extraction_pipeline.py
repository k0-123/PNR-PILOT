import pytest
from openpyxl import load_workbook

from app.core.models import ExtractionStatus, ImageStatus, JobStatus
from app.core.storage import StorageError
from app.excel.writer import EXTRACTED_COLUMNS
from app.extraction.pipeline import add_images, extract_job
from tests.conftest import gemini_row, png


@pytest.fixture
def job(db):
    return db.create_job("t")


def run(db, storage, settings, client, job):
    return extract_job(db, storage, settings, client, job)


def test_end_to_end_writes_excel_1(db, storage, settings, make_client, job):
    add_images(db, storage, settings, job, [("a.png", png(1)), ("b.png", png(2))])
    client, fake = make_client([
        {"rows": [gemini_row(),
                  gemini_row(line_no="089", surname="AKULA", first_name="MALINI MS", pnr="8SQLQK",
                             pnr_confidence=0.6, notes="glare over PNR")]},
        {"rows": [gemini_row(line_no="090", surname="=HYPERLINK(1)", pnr="AAAAAA", office_id=None)]},
    ])

    s = run(db, storage, settings, client, job)

    assert (s.images_extracted, s.images_failed, s.rows_total) == (2, 0, 3)
    assert (s.rows_ready, s.rows_review) == (1, 2)
    assert s.status == JobStatus.AWAITING_REVIEW
    assert s.cost_usd == pytest.approx(2 * 0.0008)
    assert db.get_job(job)["total_cost_usd"] == pytest.approx(s.cost_usd)

    path = storage.path(s.excel_key)
    assert path.name == f"job_{job}_extracted.xlsx"
    ws = load_workbook(path).active
    assert [c.value for c in ws[1]] == [h for h, _ in EXTRACTED_COLUMNS]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    col = {h: i for i, (h, _) in enumerate(EXTRACTED_COLUMNS)}
    assert rows[0][col["Extraction Status"]] == "READY"
    assert (rows[1][col["First Name"]], rows[1][col["Title"]]) == ("MALINI", "MS")
    assert "glare over PNR" in rows[1][col["Notes"]]
    cell = ws.cell(row=4, column=col["Surname"] + 1)
    assert (cell.value, cell.data_type, cell.quotePrefix) == ("=HYPERLINK(1)", "s", True)  # text, not a formula
    assert rows[2][col["Office ID"]] is None             # null -> empty cell


def test_all_ready_job_is_ready_for_lookup(db, storage, settings, make_client, job):
    add_images(db, storage, settings, job, [("a.png", png())])
    client, _ = make_client([{"rows": [gemini_row()]}])
    assert run(db, storage, settings, client, job).status == JobStatus.READY_FOR_LOOKUP


def test_empty_rows_job_fails(db, storage, settings, make_client, job):
    add_images(db, storage, settings, job, [("a.png", png())])
    client, _ = make_client([{"rows": []}])
    s = run(db, storage, settings, client, job)
    assert s.status == JobStatus.FAILED and s.excel_key is None
    assert "no passenger rows" in db.get_job(job)["error_message"]


def test_failed_image_marks_image_not_job_and_resumes(db, storage, settings, make_client, job):
    add_images(db, storage, settings, job, [("a.png", png(1)), ("b.png", png(2))])
    client, _ = make_client(["bad", "bad", {"rows": [gemini_row()]}])  # image a: parse fails twice
    s = run(db, storage, settings, client, job)
    assert (s.images_extracted, s.images_failed, s.status) == (1, 1, JobStatus.READY_FOR_LOOKUP)
    failed = db.get_images(job)[0]
    assert failed["status"] == ImageStatus.EXTRACTION_FAILED
    assert failed["raw_ai_response"] == "bad" and failed["attempts"] == 2

    # rerun: only the failed image goes to Gemini
    client2, fake2 = make_client([{"rows": [gemini_row(pnr="ZZZZZ9")]}])
    s2 = run(db, storage, settings, client2, job)
    assert len(fake2.calls) == 1 and (s2.images_extracted, s2.images_failed) == (2, 0)


def test_same_image_never_sent_twice(db, storage, settings, make_client):
    job1, job2 = db.create_job("one"), db.create_job("two")
    add_images(db, storage, settings, job1, [("a.png", png(7))])
    client, fake = make_client([{"rows": [gemini_row()]}])
    run(db, storage, settings, client, job1)
    run(db, storage, settings, client, job1)  # re-run: nothing pending
    add_images(db, storage, settings, job2, [("copy.png", png(7))])  # same bytes, new job
    s = run(db, storage, settings, client, job2)
    assert len(fake.calls) == 1
    assert s.rows_total == 1 and db.get_images(job2)[0]["from_cache"] == 1
    assert db.job_cost(job2)["cost_usd"] == 0


def test_duplicates_are_surname_plus_pnr(db, storage, settings, make_client, job):
    add_images(db, storage, settings, job, [("a.png", png())])
    client, _ = make_client([{"rows": [
        gemini_row(line_no="098", surname="HARIDOSS", pnr="9P4BIE"),
        gemini_row(line_no="101", surname="KISHTAPPA NAIDU", pnr="9P4BIE"),  # same PNR, other pax
        gemini_row(line_no="102", surname="HARIDOSS", pnr="9P4BIE"),         # real duplicate
    ]}])
    s = run(db, storage, settings, client, job)
    rows = db.get_rows(job)
    assert [r["extraction_status"] for r in rows] == ["READY", "READY", "DUPLICATE"]
    assert rows[2]["duplicate_of"] == rows[0]["id"]
    assert s.rows_duplicate == 1


def test_review_edit_keeps_original_and_rechecks_duplicates(db, storage, settings, make_client, job):
    add_images(db, storage, settings, job, [("a.png", png())])
    client, _ = make_client([{"rows": [gemini_row(), gemini_row(line_no="091", pnr="BOGZG0",
                                                                pnr_confidence=0.4)]}])
    run(db, storage, settings, client, job)
    second = db.get_rows(job)[1]

    db.review_row(second["id"], surname="SHUKLA", first_name=None, pnr="BOGZGQ",
                  status=ExtractionStatus.APPROVED)
    r = db.get_rows(job)[1]
    assert r["extraction_status"] == "DUPLICATE"
    assert (r["pnr"], r["edited_pnr"], r["effective_pnr"]) == ("BOGZG0", "BOGZGQ", "BOGZGQ")

    db.review_row(second["id"], surname="SHUKLA", first_name=None, pnr="BOGZGA",
                  status=ExtractionStatus.APPROVED)
    assert db.get_rows(job)[1]["extraction_status"] == "APPROVED"


def test_missing_client_fails_image_not_crash(db, storage, settings, job):
    add_images(db, storage, settings, job, [("a.png", png(9))])
    s = run(db, storage, settings, None, job)
    assert s.images_failed == 1 and s.status == JobStatus.FAILED


# ---- upload checks ----

@pytest.mark.parametrize("name,data,msg", [
    ("notes.txt", b"hello", "only PNG"),
    ("fake.png", b"GIF89a....", "not a valid"),
    ("empty.jpg", b"", "empty"),
])
def test_bad_uploads_rejected(db, storage, settings, job, name, data, msg):
    with pytest.raises(StorageError, match=msg):
        add_images(db, storage, settings, job, [(name, data)])
    assert db.count_images(job) == 0


def test_upload_limits(db, storage, settings, job):
    settings.max_file_size_mb = 0.00001
    with pytest.raises(StorageError, match="larger than"):
        add_images(db, storage, settings, job, [("a.png", png() * 100)])
    settings.max_file_size_mb, settings.max_images_per_job = 20, 2
    with pytest.raises(StorageError, match="at most 2"):
        add_images(db, storage, settings, job, [(f"{i}.png", png(i)) for i in range(3)])


def test_same_file_twice_in_job_ignored(db, storage, settings, job):
    add_images(db, storage, settings, job, [("a.png", png()), ("b.png", png())])
    assert db.count_images(job) == 1


def test_uploaded_filename_is_sanitised(db, storage, settings, job):
    add_images(db, storage, settings, job, [("..\\..\\evil name?.png", png())])
    img = db.get_images(job)[0]
    assert img["filename"] == "evil_name_.png"
    assert storage.path(img["stored_path"]).is_relative_to(storage.root / "uploads")


def test_family_members_on_one_booking_are_not_duplicates(db, storage, settings, make_client, job):
    add_images(db, storage, settings, job, [("a.png", png())])
    client, _ = make_client([{"rows": [
        gemini_row(line_no="001", surname="CHEN", first_name="LIJING", pnr="ER7P5B"),
        gemini_row(line_no="002", surname="CHEN", first_name="WEI", pnr="ER7P5B"),     # family member
        gemini_row(line_no="003", surname="CHEN", first_name="LIJING", pnr="ER7P5B"),  # same person again
    ]}])
    run(db, storage, settings, client, job)
    assert [r["extraction_status"] for r in db.get_rows(job)] == ["READY", "READY", "DUPLICATE"]
